from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.utils import (
    get_column_letter,
)
from pattern_engine import pattern_sheet_specs


HEADER_FILL = PatternFill(
    "solid",
    fgColor="D9EAF7",
)

ACTION_FILLS = {
    "BUY NOW": PatternFill(
        "solid",
        fgColor="C6EFCE",
    ),
    "WATCH": PatternFill(
        "solid",
        fgColor="FFF2CC",
    ),
    "BUY ON BREAKOUT": PatternFill(
        "solid",
        fgColor="DDEBF7",
    ),
    "BREAKOUT / VOLUME PENDING": PatternFill(
        "solid",
        fgColor="FCE4D6",
    ),
    "EXTENDED / DO NOT CHASE": PatternFill(
        "solid",
        fgColor="F4CCCC",
    ),
    "AVOID": PatternFill(
        "solid",
        fgColor="D9D9D9",
    ),
}


def format_worksheet(
    worksheet,
) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    for cell in worksheet[1]:
        cell.font = Font(
            bold=True
        )
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    action_column = None

    for cell in worksheet[1]:
        if cell.value == "Action":
            action_column = cell.column
            break

    if action_column is not None:
        for row_index in range(
            2,
            worksheet.max_row + 1,
        ):
            action_cell = worksheet.cell(
                row=row_index,
                column=action_column,
            )
            fill = ACTION_FILLS.get(
                str(action_cell.value)
            )

            if fill:
                action_cell.fill = fill
                action_cell.font = Font(
                    bold=True
                )

    for index, cells in enumerate(
        worksheet.columns,
        start=1,
    ):
        max_length = max(
            len(
                ""
                if cell.value is None
                else str(cell.value)
            )
            for cell in cells
        )

        worksheet.column_dimensions[
            get_column_letter(index)
        ].width = min(
            max_length + 2,
            40,
        )


def create_excel_report(
    report: pd.DataFrame,
    sector_summary: pd.DataFrame,
    market_status: dict[str, object],
    errors: list[dict[str, str]],
    output_file: Path,
    watchlist: pd.DataFrame,
) -> None:
    decision_columns = [
        column
        for column in [
            "Rank",
            "Symbol",
            "Company Name",
            "Industry",
            "Qualified Scanner",
            "Breakout Status",
            "Breakout Score",
            "Breakout Scanner Eligible",
            "Daily Gain %",
            "5D Return %",
            "Historical Volume Ratio",
            "Breakout Date",
            "Breakout Days Ago",
            "Primary Pattern",
            "Matched Patterns",
            "Pattern Count",
            "Pattern Confidence",
            "Pattern Pivot Price",
            "Multi-Pattern Confirmation",
            "Fundamentals Gate Status",
            "Fundamentals Pass",
            "EPS Growth YoY %",
            "Previous EPS Growth YoY %",
            "Sales Growth YoY %",
            "Operating Margin %",
            "Previous Operating Margin %",
            "RS Line New High",
            "RS Line Leads Price",
            "RS Line 20D Slope %",
            "RS Line Status",
            "Base Stage",
            "Base Stage Pass",
            "Current Price",
            "Current Price Source",
            "Action",
            "Entry Window",
            "Live Extension %",
            "Confidence Score",
            "Trade Urgency",
            "Risk Level",
            "Trade Quality Score",
            "Trade Grade",
            "Suggested Entry",
            "Suggested Stop Loss",
            "Stop Loss %",
            "Suggested Quantity",
            "Planned Investment",
            "Actual Rupee Risk",
            "Target 1:2",
            "Target 1:3",
            "EMA 10",
            "EMA 20",
            "EMA 50",
            "EMA Bullish Alignment",
            "RS Score",
            "Sector RS Score",
            "Volume Score",
            "VCP Score",
            "Decision Reason",
            "Why Selected",
        ]
        if column in report.columns
    ]

    trade_planner_columns = [
        column
        for column in [
            "Rank",
            "Symbol",
            "Current Price",
            "Suggested Entry",
            "Entry Distance %",
            "Live Extension %",
            "Suggested Stop Loss",
            "Stop Loss %",
            "Risk Per Share",
            "Suggested Quantity",
            "Planned Investment",
            "Actual Rupee Risk",
            "1R Price",
            "Target 1:2",
            "Target 1:3",
            "Breakeven Stop",
            "ATR 14",
            "ATR % of Entry",
            "Recent 10D Low",
            "EMA 10",
            "EMA 20",
            "EMA 50",
            "EMA Bullish Alignment",
            "Trailing Stop Plan",
            "Trade Plan Valid",
        ]
        if column in report.columns
    ]

    pattern_summary_rows = []
    for pattern_name, match_column in pattern_sheet_specs():
        if match_column in report.columns:
            matches = report[match_column].fillna(False).astype(bool)
            pattern_summary_rows.append(
                {
                    "Pattern": pattern_name,
                    "Matches": int(matches.sum()),
                    "Near Trigger": int(
                        (
                            matches
                            & report["Action"].eq("NEAR TRIGGER")
                        ).sum()
                    ),
                    "Buy Now": int(
                        (
                            matches
                            & report["Action"].eq("BUY NOW")
                        ).sum()
                    ),
                }
            )

    sheets: dict[
        str,
        pd.DataFrame,
    ] = {
        "Decision Board": report[
            decision_columns
        ],
        "Buy Now": report[
            report["Action"].eq(
                "BUY NOW"
            )
        ],
        "Watch": report[
            report["Action"].isin(
                [
                    "WATCH",
                    "BUY ON BREAKOUT",
                ]
            )
        ],
        "Hold-No Add": report[
            report["Action"].eq(
                "BREAKOUT / VOLUME PENDING"
            )
        ],
        "Extended": report[
            report["Action"].eq(
                "EXTENDED / DO NOT CHASE"
            )
        ],
        "All Stocks": report,
        "Top Quality Trades": report[
            report["Top Quality Trade"]
        ],
        "Trade Ready": report[
            report["Trade Ready"]
        ],
        "Trade Planner": report[
            trade_planner_columns
        ],
        "Passing Stocks": report[
            report[
                "Trend Template Pass"
            ]
        ],
        "RS Leaders": report[
            report["RS Pass"]
        ],
        "Sector Leaders": (
            sector_summary
        ),
        "Volume Leaders": report[
            report[
                "Volume Setup Pass"
            ]
        ],
        "VCP Candidates": report[
            report["VCP Pass"]
        ],
        "Final Candidates": report[
            report[
                "Trend + RS + Volume + VCP Pass"
            ]
        ],
        "Near Pivot": report[
            report[
                "Ready Near Pivot"
            ]
        ],
        "Breakouts": report[
            report.get(
                "Breakout Status",
                pd.Series("", index=report.index),
            ).eq("BREAKOUT TODAY")
        ],
        "Recent Breakouts": report[
            report.get(
                "Breakout Status",
                pd.Series("", index=report.index),
            ).eq("RECENT BREAKOUT")
        ],
        "Approaching": report[
            report.get(
                "Breakout Status",
                pd.Series("", index=report.index),
            ).eq("APPROACHING BREAKOUT")
        ],
        "Momentum Movers": report[
            report.get(
                "Breakout Status",
                pd.Series("", index=report.index),
            ).eq("MOMENTUM MOVER")
        ],
        "Daily Watchlist": (
            watchlist
        ),
        "Pattern Summary": pd.DataFrame(
            pattern_summary_rows
        ),
        "Market Status": pd.DataFrame(
            [market_status]
        ),
        "Errors": (
            pd.DataFrame(
                errors
            )
            if errors
            else pd.DataFrame(
                columns=[
                    "Symbol",
                    "Data Source",
                    "Data Symbol",
                    "Company Name",
                    "Error",
                ]
            )
        ),
    }

    for pattern_name, match_column in pattern_sheet_specs():
        if match_column in report.columns:
            sheets[pattern_name] = report[
                report[match_column].fillna(False).astype(bool)
            ]

    output_file.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        output_file,
        engine="openpyxl",
    ) as writer:
        for (
            sheet_name,
            dataframe,
        ) in sheets.items():
            dataframe.to_excel(
                writer,
                sheet_name=sheet_name[
                    :31
                ],
                index=False,
            )

        for worksheet in (
            writer.sheets.values()
        ):
            format_worksheet(
                worksheet
            )
